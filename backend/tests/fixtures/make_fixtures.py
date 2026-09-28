"""
Regenerate the binary resume fixtures from sample_resume.txt:

    python -m tests.fixtures.make_fixtures      (run from backend/)

- sample_resume.pdf  — a real single-page text PDF, written by hand (no PDF
  library needed) using the built-in Helvetica font with WinAnsi encoding so
  bullets (•) and dashes (–, —) survive extraction.
- sample_resume.docx — built with python-docx; bullets use the "List Bullet"
  style (no literal "•" in the text), exercising the parser's list detection.

Both files are committed so tests (and Swagger demos) don't need to run this.
"""

from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "sample_resume.txt"


def _pdf_escape(text: str) -> bytes:
    raw = text.encode("cp1252", errors="replace")
    return raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def make_pdf(lines: list[str]) -> bytes:
    """Minimal valid PDF 1.4: one page, one Helvetica text stream, correct xref table."""
    content = [b"BT", b"/F1 10 Tf", b"13 TL", b"50 760 Td"]
    for line in lines:
        content.append(b"(" + _pdf_escape(line) + b") Tj T*")
    content.append(b"ET")
    stream = b"\n".join(content)

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]

    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    return bytes(out)


def make_docx(lines: list[str]) -> bytes:
    import io

    from docx import Document

    doc = Document()
    for line in lines:
        if line.startswith("• "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.isupper() and line.strip():
            doc.add_heading(line.title(), level=2)
        else:
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def main() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    (HERE / "sample_resume.pdf").write_bytes(make_pdf(lines))
    (HERE / "sample_resume.docx").write_bytes(make_docx(lines))
    print("wrote sample_resume.pdf and sample_resume.docx")


if __name__ == "__main__":
    main()
